import asyncio
import os
import sys
from playwright.async_api import async_playwright
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r'c:\Users\Giang\Desktop\slide_new'
EXPORTS_DIR = os.path.join(BASE_DIR, 'exports')
SLIDES_DIR = os.path.join(EXPORTS_DIR, 'slides')

os.makedirs(SLIDES_DIR, exist_ok=True)

PPTX_ROOT_PATH = os.path.join(BASE_DIR, 'Thuyet_Trinh_QPAN_DaiNam_2K.pptx')
PDF_ROOT_PATH = os.path.join(BASE_DIR, 'Thuyet_Trinh_QPAN_DaiNam_2K.pdf')

PPTX_EXPORT_PATH = os.path.join(EXPORTS_DIR, 'Thuyet_Trinh_QPAN_DaiNam_2K.pptx')
PDF_EXPORT_PATH = os.path.join(EXPORTS_DIR, 'Thuyet_Trinh_QPAN_DaiNam_2K.pdf')

TOTAL_SLIDES = 10

async def capture_all_slides():
    print("Bắt đầu mở trình duyệt Chrome chụp 10 Slide ở chuẩn Canvas 2K (2560x1440)...")
    image_paths = []

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            executable_path=r'C:\Program Files\Google\Chrome\Application\chrome.exe',
            headless=True
        )
        context = await browser.new_context(
            viewport={'width': 1920, 'height': 1080},
            device_scale_factor=2560 / 1920
        )
        page = await context.new_page()
        file_url = os.path.join(BASE_DIR, 'index.html').replace('\\', '/')
        await page.goto(f'file:///{file_url}')
        await page.wait_for_timeout(1000)

        # Apply exporting-2k class to strip transforms, transitions, web UI and timers
        await page.evaluate("document.body.classList.add('exporting-2k')")
        await page.wait_for_timeout(200)

        for i in range(1, TOTAL_SLIDES + 1):
            slide_id = f"slide-{i}"
            img_filename = f"slide_{i:02d}.png"
            img_path = os.path.join(SLIDES_DIR, img_filename)

            # Switch slide
            await page.evaluate(f"""() => {{
                document.querySelectorAll('.slide').forEach(s => s.classList.remove('active'));
                const el = document.getElementById('{slide_id}');
                if (el) el.classList.add('active');
                if (window.lucide) window.lucide.createIcons();
            }}""")

            await page.wait_for_timeout(250)

            target_locator = page.locator(f"#{slide_id}")
            await target_locator.screenshot(path=img_path)

            im = Image.open(img_path)
            print(f"  [OK] Đã chụp Slide {i:02d}/10 -> {img_filename} (kích thước {im.size[0]}x{im.size[1]})")
            image_paths.append(img_path)

        await browser.close()

    return image_paths

def build_pptx(image_paths):
    print("Đang đóng gói file PowerPoint (.pptx)...")
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    for img_path in image_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            img_path,
            Inches(0), Inches(0),
            width=prs.slide_width,
            height=prs.slide_height
        )

    prs.save(PPTX_ROOT_PATH)
    prs.save(PPTX_EXPORT_PATH)
    print(f"  [OK] Đã lưu file PPTX: {PPTX_ROOT_PATH} ({os.path.getsize(PPTX_ROOT_PATH):,} bytes)")

def build_pdf(image_paths):
    print("Đang đóng gói file PDF (2K Landscape)...")
    pil_images = []
    for path in image_paths:
        im = Image.open(path).convert('RGB')
        pil_images.append(im)

    if pil_images:
        pil_images[0].save(
            PDF_ROOT_PATH,
            save_all=True,
            append_images=pil_images[1:],
            resolution=150.0
        )
        pil_images[0].save(
            PDF_EXPORT_PATH,
            save_all=True,
            append_images=pil_images[1:],
            resolution=150.0
        )
        print(f"  [OK] Đã lưu file PDF: {PDF_ROOT_PATH} ({os.path.getsize(PDF_ROOT_PATH):,} bytes)")

async def main():
    image_paths = await capture_all_slides()
    build_pptx(image_paths)
    build_pdf(image_paths)
    print("HOÀN THÀNH XUẤT TOÀN BỘ 10 SLIDE CANVAS 2K THÀNH PPTX VÀ PDF!")

if __name__ == '__main__':
    asyncio.run(main())
